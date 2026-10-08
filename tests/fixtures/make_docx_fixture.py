"""Regenerate sample.docx: python3 -m pip install python-docx && python3 make_docx_fixture.py

Covers what the docx reader/writer must round-trip: styles, run formatting,
hyperlinks, bookmarks, fields, nested lists, tabs, page/section breaks, images,
merged + shaded table cells and a content control (sdt).
"""
import io
import struct
import zlib

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Inches


def png(w=4, h=4, rgb=(0, 0, 255)):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


d = Document()

# 0: heading
d.add_heading("Doc Heading", level=1)

# 1: run formatting
p = d.add_paragraph("Plain ")
p.add_run("bold").bold = True
p.add_run(" and ")
p.add_run("italic").italic = True
r = p.add_run(" red20")
r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
r.font.size = Pt(20)
r = p.add_run(" yellow")
r.font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW
p.add_run("\tafter tab")

# 2: hyperlink + bookmark + simple field, all in one paragraph
p = d.add_paragraph()
p._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="0" w:name="mark"/>'))
p.add_run("See ")
rid = d.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
p._p.append(parse_xml(
    f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}" w:history="1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>example</w:t></w:r></w:hyperlink>'))
p._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="0"/>'))
p.add_run(" page ")
p._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))

# 3-5: list with a nested level
d.add_paragraph("First bullet", style="List Bullet")
d.add_paragraph("Nested bullet", style="List Bullet 2")
d.add_paragraph("Numbered", style="List Number")

# 6: page break inside a paragraph, plus a rendering hint that must not become a break
p = d.add_paragraph("Before break")
p.add_run().add_break(WD_BREAK.PAGE)
p.add_run("after break")
p.runs[0]._r.append(parse_xml(f'<w:lastRenderedPageBreak {nsdecls("w")}/>'))

# 7: image
d.add_paragraph().add_run().add_picture(io.BytesIO(png()), width=Inches(1))

# 8: table with merged + shaded cells
t = d.add_table(rows=2, cols=3)
t.style = "Table Grid"
t.cell(0, 0).merge(t.cell(0, 1))
t.cell(0, 0).text = "Merged"
t.cell(0, 2).text = "C"
t.cell(1, 0).text = "Shaded"
tcPr = t.cell(1, 0)._tc.get_or_add_tcPr()
tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="FFFF00"/>'))

# 9: paragraph that ends a section (sectPr inside pPr)
p = d.add_paragraph("End of section one")
p._p.get_or_add_pPr().append(parse_xml(
    f'<w:sectPr {nsdecls("w")}><w:type w:val="continuous"/><w:pgSz w:w="12240" w:h="15840"/></w:sectPr>'))

# 10: content control (block sdt)
d.element.body.insert(len(d.element.body) - 1, parse_xml(
    f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Box"/></w:sdtPr><w:sdtContent>'
    f'<w:p><w:r><w:t>Inside control</w:t></w:r></w:p></w:sdtContent></w:sdt>'))

# 11: symbol run (opaque run content) mid-paragraph
p = d.add_paragraph("Sym ")
p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:sym w:font="Wingdings" w:char="F04A"/></w:r>'))
p.add_run(" end")

# 12: list paragraph with direct numbering at level 1
p = d.add_paragraph("Direct nested")
p._p.get_or_add_pPr().append(parse_xml(
    f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>'))

d.save("sample.docx")
